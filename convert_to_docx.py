import re
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Paths
INPUT_MD = r"\\wsl.localhost\Ubuntu\home\shery\projects\background-job-service\SYSTEM_DESIGN_AND_CONCEPTS.md"
OUTPUT_DOCX = r"\\wsl.localhost\Ubuntu\home\shery\projects\background-job-service\SYSTEM_DESIGN_AND_CONCEPTS.docx"

# Color Palette
COLOR_NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # #1E3A8A - H1 & Table Headers
COLOR_SLATE_BLUE = RGBColor(0x25, 0x63, 0xEB) # #2563EB - H2 & Accents
COLOR_DARK_GRAY = RGBColor(0x37, 0x41, 0x51)  # #374151 - H3 & Body
COLOR_BODY_TEXT = RGBColor(0x1F, 0x29, 0x37)  # #1F2937 - Body Text
COLOR_MUTED_TEXT = RGBColor(0x6B, 0x72, 0x80) # #6B7280 - Subtitles / Footnotes
COLOR_CODE_TEXT = RGBColor(0x11, 0x18, 0x27)  # #111827 - Monospace code
COLOR_WARNING_BORDER = RGBColor(0xEF, 0x44, 0x44)

HEX_NAVY = "1E3A8A"
HEX_SLATE_BLUE = "2563EB"
HEX_CODE_BG = "F3F4F6"
HEX_CODE_BORDER = "E5E7EB"
HEX_ZEBRA_ROW = "F9FAFB"
HEX_TABLE_BORDER = "D1D5DB"
HEX_CALLOUT_BLUE_BG = "EFF6FF"
HEX_CALLOUT_BLUE_BORDER = "2563EB"
HEX_CALLOUT_WARN_BG = "FEF2F2"
HEX_CALLOUT_WARN_BORDER = "DC2626"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    # Measurements in dxa (1 pt = 20 dxa)
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_callout_borders(cell, border_hex="2563EB"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_code_block_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_SLATE_BLUE}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_CODE_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_table_cell_borders(cell, border_hex="D1D5DB"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{border_hex}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_inline_formatted_text(paragraph, text, default_font="Calibri", default_size=Pt(11), default_color=COLOR_BODY_TEXT):
    # Regex to split on bold, italic, inline code, and links
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))')
    tokens = pattern.split(text)

    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = default_font
            run.font.size = default_size
            run.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = default_font
            run.font.size = default_size
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x9D, 0x17, 0x4D) # Deep magenta-rose for inline code
            run.bold = True
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                link_text, link_url = m.groups()
                run = paragraph.add_run(link_text)
                run.font.name = default_font
                run.font.size = default_size
                run.font.color.rgb = COLOR_SLATE_BLUE
                run.underline = True
            else:
                run = paragraph.add_run(token)
                run.font.name = default_font
                run.font.size = default_size
                run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(token)
            run.font.name = default_font
            run.font.size = default_size
            run.font.color.rgb = default_color

def create_callout_box(doc, text_lines, is_warning=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_hex = HEX_CALLOUT_WARN_BG if is_warning else HEX_CALLOUT_BLUE_BG
    border_hex = HEX_CALLOUT_WARN_BORDER if is_warning else HEX_CALLOUT_BLUE_BORDER
    set_cell_background(cell, bg_hex)
    set_callout_borders(cell, border_hex)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    for i, line in enumerate(text_lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
        
        cleaned = re.sub(r'^[>\s]+', '', line)
        if is_warning and i == 0 and not cleaned.startswith('**WARNING'):
            r = p.add_run("WARNING: ")
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_WARNING_BORDER

        add_inline_formatted_text(p, cleaned, default_size=Pt(10.5))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_code_block(doc, code_lines, language=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)

    set_cell_background(cell, HEX_CODE_BG)
    set_code_block_borders(cell)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05

    for i, line in enumerate(code_lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
        
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_CODE_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_markdown_table(doc, table_lines):
    parsed_rows = []
    for line in table_lines:
        line = line.strip()
        if not line or not line.startswith('|'):
            continue
        # Split by pipe and strip
        cols = [c.strip() for c in line.split('|')[1:-1]]
        # Skip markdown separator row |---|---|
        if cols and all(re.match(r'^:?-+:?$', c) for c in cols):
            continue
        parsed_rows.append(cols)

    if not parsed_rows:
        return

    num_rows = len(parsed_rows)
    num_cols = max(len(r) for r in parsed_rows)

    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True

    # Style Header Row
    header_cols = parsed_rows[0]
    for c_idx in range(num_cols):
        cell = tbl.cell(0, c_idx)
        val = header_cols[c_idx] if c_idx < len(header_cols) else ""
        set_cell_background(cell, HEX_NAVY)
        set_table_cell_borders(cell, HEX_NAVY)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        run = p.add_run(val)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Style Data Rows
    for r_idx in range(1, num_rows):
        row_data = parsed_rows[r_idx]
        bg_hex = HEX_ZEBRA_ROW if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx in range(num_cols):
            cell = tbl.cell(r_idx, c_idx)
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            set_cell_background(cell, bg_hex)
            set_table_cell_borders(cell, HEX_TABLE_BORDER)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            
            # Format inline elements (like bold / code)
            add_inline_formatted_text(p, val, default_size=Pt(9.5))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def convert():
    print(f"Reading Markdown: {INPUT_MD}")
    with open(INPUT_MD, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Configure Margins: 1 inch (72pt) all around
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set normal style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = COLOR_BODY_TEXT

    in_code_block = False
    code_block_lines = []
    code_block_lang = ""

    in_table = False
    table_lines = []

    in_callout = False
    callout_lines = []
    callout_is_warning = False

    for idx, raw_line in enumerate(lines):
        line = raw_line.rstrip('\r\n')

        # 1. Handle Code Blocks
        if line.startswith('```'):
            if in_code_block:
                create_code_block(doc, code_block_lines, code_block_lang)
                code_block_lines = []
                code_block_lang = ""
                in_code_block = False
            else:
                # Flush pending table or callout
                if in_table:
                    create_markdown_table(doc, table_lines)
                    table_lines = []
                    in_table = False
                if in_callout:
                    create_callout_box(doc, callout_lines, callout_is_warning)
                    callout_lines = []
                    in_callout = False

                in_code_block = True
                code_block_lang = line[3:].strip()
            continue

        if in_code_block:
            code_block_lines.append(line)
            continue

        # 2. Handle Tables
        if line.strip().startswith('|'):
            if not in_table:
                if in_callout:
                    create_callout_box(doc, callout_lines, callout_is_warning)
                    callout_lines = []
                    in_callout = False
                in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            create_markdown_table(doc, table_lines)
            table_lines = []
            in_table = False

        # 3. Handle Callouts & Quotes
        if line.strip().startswith('>'):
            if not in_callout:
                in_callout = True
                callout_is_warning = False
            if '[!WARNING]' in line or '[!CAUTION]' in line:
                callout_is_warning = True
            callout_lines.append(line)
            continue
        elif in_callout:
            create_callout_box(doc, callout_lines, callout_is_warning)
            callout_lines = []
            in_callout = False

        # 4. Empty Lines
        if not line.strip():
            continue

        # 5. Section Divider
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("―" * 40)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xDC, 0xDE, 0xE2)
            continue

        # 6. Document Title: # Title
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(title_text)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(24)
            run.font.color.rgb = COLOR_NAVY
            continue

        # 7. Heading 1: ## Section
        if line.startswith('## '):
            h1_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(22)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h1_text)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.font.color.rgb = COLOR_NAVY
            continue

        # 8. Heading 2: ### Subsection
        if line.startswith('### '):
            h2_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.color.rgb = COLOR_SLATE_BLUE
            continue

        # 9. Heading 3: #### Minor Header
        if line.startswith('#### '):
            h3_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h3_text)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_DARK_GRAY
            continue

        # 10. Bullet Lists: - or *
        if re.match(r'^\s*[-*]\s+', line):
            indent_level = len(re.match(r'^\s*', line).group(0)) // 2
            bullet_text = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            add_inline_formatted_text(p, bullet_text)
            continue

        # 11. Numbered Lists: 1. , 2.
        if re.match(r'^\s*\d+\.\s+', line):
            num_match = re.match(r'^\s*(\d+\.)\s+(.*)', line)
            prefix = num_match.group(1)
            content_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.3)
            r_num = p.add_run(f"{prefix} ")
            r_num.bold = True
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(11)
            r_num.font.color.rgb = COLOR_NAVY
            add_inline_formatted_text(p, content_text)
            continue

        # 12. Regular Paragraph Body Text
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_formatted_text(p, line)

    # Flush any trailing structures
    if in_table:
        create_markdown_table(doc, table_lines)
    if in_callout:
        create_callout_box(doc, callout_lines, callout_is_warning)
    if in_code_block:
        create_code_block(doc, code_block_lines, code_block_lang)

    print(f"Saving formatted Word document to: {OUTPUT_DOCX}")
    doc.save(OUTPUT_DOCX)
    size_bytes = os.path.getsize(OUTPUT_DOCX)
    print(f"SUCCESS: Generated Word Document successfully! Size: {size_bytes / 1024:.1f} KB")

if __name__ == "__main__":
    convert()
